"""B2 whole-file attacks: empty, header-only, binary, wrong extension, XLSX, JSON, scale."""
import json
import os
import random
import time
from datetime import date, datetime

import openpyxl

from tests.breakers.edge_world import HR, HR_H, LIC, NOISE, PAY, run

IDS = ["P-E201", "P-E202", "P-E203", "P-E204"]


def _ids(r):
    return [p["person_id"] for p in r.persons]


def test_empty_file_is_quarantined_not_crashed(tmp_path, settings):
    empty = tmp_path / "empty.csv"
    empty.write_text("")
    r = run(tmp_path, settings, extra=[empty])
    assert _ids(r) == IDS and [i.severity for i in r.ids("FILE-QUARANTINED")] == ["HIGH"]


def test_header_only_file_is_flagged(tmp_path, settings):
    f = tmp_path / "hdr.csv"
    f.write_text(",".join(HR_H) + "\n")
    r = run(tmp_path, settings, extra=[f])
    assert r.attention() - NOISE


def test_binary_garbage_named_csv_is_quarantined(tmp_path, settings):
    f = tmp_path / "garbage.csv"
    f.write_bytes(os.urandom(5000))
    r = run(tmp_path, settings, extra=[f])
    assert _ids(r) == IDS and r.ids("FILE-QUARANTINED")


def test_csv_named_txt_is_read(tmp_path, settings):
    f = tmp_path / "hr_export.txt"
    f.write_text(",".join(HR_H) + "\n" + "\n".join(",".join(r) for r in HR) + "\n")
    r = run(tmp_path, settings, hr=None, extra=[f])
    assert _ids(r) == IDS


def test_json_export_of_the_roster(tmp_path, settings):
    f = tmp_path / "roster.json"
    f.write_text(json.dumps({"employees": [dict(zip(HR_H, r)) for r in HR]}))
    r = run(tmp_path, settings, hr=None, extra=[f])
    assert _ids(r) == IDS


def test_xlsx_with_roster_sheet_and_notes_sheet(tmp_path, settings):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Roster"
    ws.append(HR_H)
    for row in HR:
        ws.append(row)
    wb.create_sheet("Notes").append(["Roster exported 2026-09-20; ignore terminated staff"])
    f = tmp_path / "roster.xlsx"
    wb.save(f)
    r = run(tmp_path, settings, hr=None, extra=[f])
    assert _ids(r) == IDS
    assert not [i for i in r.ids("TABLE-UNMAPPED") if i.severity in ("HIGH", "CRITICAL")]
    assert not r.ids("AGENT-UNAVAILABLE"), "a free-text Notes sheet should not need a schema-mapper agent"


def test_xlsx_native_dates_numeric_ids_and_title_row(tmp_path, settings):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Harborview HR export"])
    ws.append([])
    ws.append(HR_H)
    for i, row in enumerate(HR):
        ws.append([201 + i, *row[1:6], row[6], datetime(*map(int, row[7].split("-"))), datetime(*map(int, row[8].split("-")))])
    f = tmp_path / "roster.xlsx"
    wb.save(f)
    r = run(tmp_path, settings, hr=None, extra=[f])
    assert len(r.persons) == 4 and all(p["has_hr"] for p in r.persons)
    assert [c["expires_on"] for c in r.creds(r.persons[0]["person_id"])] == [date(2027, 5, 31)]


def _synthetic_world(n):
    rng = random.Random(7)
    syl = ["ka", "ri", "mo", "ta", "len", "dar", "vi", "sho", "ne", "bru", "lo", "fen", "qua", "zi", "mar", "th"]
    names = set()
    while len(names) < n:
        names.add((("".join(rng.choices(syl, k=3))).title(), ("".join(rng.choices(syl, k=4))).title()))
    hr, pay, lic = [], [], []
    titles = {"RN": "Registered Nurse", "LPN": "Licensed Practical Nurse", "CNA": "Certified Nursing Assistant"}
    for i, (f, l) in enumerate(sorted(names)):
        role = ("RN", "LPN", "CNA")[i % 3]
        fac = ("Harborview Bayside", "Harborview Riverdale")[i % 2]
        hr.append([f"E{1000 + i}", f, l, titles[role], fac, f"718-555-{i % 10000:04d}", f"{role}-{100000 + i}", "2027-05-31", "2020-03-02"])
        pay.append([f"P-{i}", f"{l.upper()}, {f.upper()}", role, ("BYS", "RVD")[i % 2], "2026-09-14", "2026-09-20", "40"])
        lic.append([f"{role}-{100000 + i}", f"{l.upper()}, {f.upper()}", role, "2027-05-31", "2026-09-01"])
    return hr, pay, lic


def test_thousand_employee_ingest_is_fast_enough(tmp_path, settings):
    """Scale probe. 50,000 rows would be 50x this; budget 30 s for 1,000 HR + 1,000 payroll + 1,000 licence rows
    (observed ~190 s => ~2.6 h for 50k)."""
    hr, pay, lic = _synthetic_world(1000)
    t0 = time.time()
    r = run(tmp_path, settings, hr=hr, pay=pay, lic=lic, no_sched=True)
    elapsed = time.time() - t0
    assert len(r.persons) == 1000
    assert elapsed < 30, f"1,000 employees took {elapsed:.0f}s"
