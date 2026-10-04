import json
import time
from datetime import date, datetime
from pathlib import Path

import openpyxl

from sot.core.models import FileRef
from sot.core.registry import ExtractContext
from sot.parsers.csv_parser import CsvParser
from sot.parsers.json_parser import JsonParser
from sot.parsers.xlsx_parser import XlsxParser
from sot.sniff.auction import run_auction

README = Path(__file__).parent / "fixtures" / "readme_sample"
PARSERS = [CsvParser(), XlsxParser(), JsonParser()]


def extract(path: Path):
    f = FileRef(file_id="a1" * 32, file_name=path.name, path=str(path), size=path.stat().st_size, received_at=datetime(2026, 1, 1))
    won = run_auction(f, PARSERS, 0.5).winner
    parser = next(p for p in PARSERS if p.name == won.parser)
    return won, parser.extract(f, ExtractContext(settings=None, pack=None, run_id="r", sniff_details=won.details))


def one(path: Path):
    _, tables = extract(path)
    assert len(tables) == 1
    return tables[0]


def test_readme_csvs():
    t = one(README / "licenses.csv")
    assert t.header == ["license_number", "name_on_license", "license_type", "expiration_date", "last_verified"]
    assert t.df["name_on_license"][0] == "REYES, SOFIA"
    assert t.df["_src_row"][0] == 2
    assert set(t.df.drop("_src_row").dtypes) == {t.df.dtypes[0]} and str(t.df.dtypes[0]) == "String"
    assert t.extraction.method == "csv"
    hr = one(README / "hr_roster.csv")
    assert hr.header[0] == "employee_id" and hr.df["employee_id"][0] == "E201" and hr.df.height >= 2
    pay = one(README / "payroll.csv")
    assert pay.df["hours_paid"][0] == "36"


def test_semicolon_and_tab(tmp_path):
    (tmp_path / "s.csv").write_text("a;b;c\n1;x,y;3\n4;5;6\n")
    (tmp_path / "t.tsv").write_text("a\tb\tc\n1\tx\t3\n4\t5\t6\n")
    w, _ = extract(tmp_path / "s.csv")
    assert w.details["delimiter"] == ";"
    assert one(tmp_path / "s.csv").df["b"].to_list() == ["x,y", "5"]
    assert extract(tmp_path / "t.tsv")[0].details["delimiter"] == "\t"
    assert one(tmp_path / "t.tsv").df["c"].to_list() == ["3", "6"]


def test_latin1_without_bom(tmp_path):
    p = tmp_path / "l.csv"
    p.write_bytes("name,city\nJosé,Bogotá\nZoë,Köln\n".encode("latin-1"))
    assert one(p).df["name"].to_list() == ["José", "Zoë"]


def test_utf8_bom(tmp_path):
    p = tmp_path / "b.csv"
    p.write_bytes("﻿name,city\nJosé,Bogotá\n".encode("utf-8"))
    t = one(p)
    assert t.header == ["name", "city"] and t.df["name"][0] == "José"


def test_title_rows_blank_row_and_src_row(tmp_path):
    p = tmp_path / "t.csv"
    p.write_text("Harborview export,,\nRun 1,,\n,,\nid,name,start\n1,Ann,2026-01-02\n2,Bob,2026-02-03\n")
    t = one(p)
    assert t.header == ["id", "name", "start"]
    assert t.df["_src_row"].to_list() == [5, 6]


def test_ragged_trailing_commas_and_blank_lines(tmp_path):
    p = tmp_path / "r.csv"
    p.write_text("id,name,\n1,Ann,\n2,Bob\n\n3,Cy,,\n\n\n")
    t = one(p)
    assert t.header == ["id", "name"]
    assert t.df["id"].to_list() == ["1", "2", "3"]
    assert t.df["_src_row"].to_list() == [2, 3, 5]


def test_multiline_quoted_field_src_row(tmp_path):
    p = tmp_path / "m.csv"
    p.write_text('id,note\n1,"a\nb"\n2,c\n')
    assert one(p).df["_src_row"].to_list() == [2, 4]


def test_xlsx_title_row_date_and_numbers(tmp_path):
    p = tmp_path / "x.csv"  # xlsx content, wrong extension
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Roster"
    ws.append(["Staff roster"])
    ws.append([])
    ws.append(["id", "hired", "hours", "name"])
    ws.append([1, date(2026, 9, 14), 36.0, "Ann"])
    ws.append([2, datetime(2026, 9, 15, 7, 30), 7.5, "Bob"])
    wb.create_sheet("Empty")
    wb.save(p)
    won, tables = extract(p)
    assert won.parser == "xlsx" and len(tables) == 1
    t = tables[0]
    assert t.sheet == "Roster" and t.table_id == "a1a1a1a1a1a1:Roster:0"
    assert t.header == ["id", "hired", "hours", "name"]
    assert t.df["id"].to_list() == ["1", "2"]
    assert t.df["hired"].to_list() == ["2026-09-14", "2026-09-15T07:30:00"]
    assert t.df["hours"].to_list() == ["36", "7.5"]
    assert t.df["_src_row"].to_list() == [4, 5]


def test_xlsx_two_sheets(tmp_path):
    p = tmp_path / "x.xlsx"
    wb = openpyxl.Workbook()
    wb.active.title = "A"
    wb.active.append(["a", "b"])
    wb.active.append(["1", "2"])
    wb.create_sheet("B").append(["c", "d"])
    wb["B"].append(["3", "4"])
    wb.save(p)
    assert [t.sheet for t in extract(p)[1]] == ["A", "B"]


def test_json_list_of_objects_and_nested(tmp_path):
    p = tmp_path / "a.json"
    p.write_text(json.dumps([{"id": 1, "p": {"x": "a"}, "t": [1, 2]}, {"id": "2", "p": {"x": "b"}, "extra": 3.0}]))
    t = one(p)
    assert t.header == ["id", "p.x", "t", "extra"]
    assert t.df["id"].to_list() == ["1", "2"]
    assert t.df["t"].to_list() == ["[1, 2]", None]
    assert t.df["extra"].to_list() == [None, "3"]
    assert t.df["_src_row"].to_list() == [1, 2]
    assert t.extraction.method == "json"


def test_json_object_of_lists_and_ndjson(tmp_path):
    p = tmp_path / "o.json"
    p.write_text(json.dumps({"note": "x", "staff": [{"a": 1}], "shifts": [{"b": 2}, {"b": 3}]}))
    _, tables = extract(p)
    assert [t.header for t in tables] == [["a"], ["b"]]
    assert len({t.table_id for t in tables}) == 2
    n = tmp_path / "n.ndjson"
    n.write_text('{"a": 1}\n\n{"a": 2, "b": "z"}\n')
    t = one(n)
    assert t.header == ["a", "b"] and t.df["_src_row"].to_list() == [1, 3]


def test_json_schedule_fixture_sniffs_json():
    won, _ = extract(README / "schedule.json")
    assert won.parser == "json"


def test_perf_200k_rows(tmp_path):
    p = tmp_path / "big.csv"
    lines = ["id,name,role,start,hours"] + [f'{i},"DOE, JOHN {i}",RN,2026-09-14,{i % 40}' for i in range(200_000)]
    p.write_text("\n".join(lines) + "\n")
    t0 = time.perf_counter()
    t = one(p)
    assert time.perf_counter() - t0 < 5
    assert t.df.height == 200_000


def test_non_western_encoding_falls_back_to_charset_normalizer(tmp_path):
    p = tmp_path / "sj.csv"
    p.write_bytes("名前,都市\n田中太郎,東京\n山田花子,大阪\n佐藤次郎,京都\n".encode("shift_jis"))
    assert one(p).df["名前"].to_list() == ["田中太郎", "山田花子", "佐藤次郎"]
